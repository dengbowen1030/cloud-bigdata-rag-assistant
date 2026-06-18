import os
import shutil
import sys
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

_tmpdir = tempfile.TemporaryDirectory()
_tmproot = Path(_tmpdir.name)
os.environ["DATABASE_URL"] = f"sqlite:///{(_tmproot / 'test_edurag.db').as_posix()}"
os.environ["UPLOAD_DIR"] = str(_tmproot / "uploads" / "raw")
os.environ["UPLOAD_PROCESSED_DIR"] = str(_tmproot / "uploads" / "processed")
os.environ["VECTOR_STORE_DIR"] = str(_tmproot / "vector_store" / "faiss_index")
os.environ["RAG_EMBEDDING_MODE"] = "mock"
os.environ["LLM_PROVIDER"] = "mock"

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas
from sqlalchemy import select

from app.core.settings import settings
from app.database import Base, SessionLocal, engine, init_db
from app.main import app
from app.models import ChunkModel
import rag.embedding as embedding_module
from rag.embedding import EmbeddingProvider
from rag.qa_chain import generate_chat_answer
from rag.retriever import Retriever


def tearDownModule():
    engine.dispose()
    _tmpdir.cleanup()


def assert_envelope(test_case, payload):
    test_case.assertEqual(set(payload.keys()), {"success", "data", "message", "error_code"})


def build_docx_bytes() -> bytes:
    buffer = BytesIO()
    document = DocxDocument()
    document.add_heading("Lifecycle Notes", level=1)
    document.add_paragraph("DOCX parsing should preserve paragraph text for chunking.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "phase"
    table.cell(0, 1).text = "testing"
    document.save(buffer)
    return buffer.getvalue()


def build_pdf_bytes() -> bytes:
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer)
    pdf.drawString(72, 720, "PDF parsing should extract visible text from page one.")
    pdf.drawString(72, 700, "This document is used for Stage 2 rebuild acceptance.")
    pdf.save()
    return buffer.getvalue()


class ApiSmokeTest(unittest.TestCase):
    def setUp(self):
        Base.metadata.drop_all(bind=engine)
        init_db()
        shutil.rmtree(settings.vector_store_dir, ignore_errors=True)
        shutil.rmtree(settings.upload_dir, ignore_errors=True)
        shutil.rmtree(settings.processed_dir, ignore_errors=True)

    def test_health_uses_unified_response(self):
        with TestClient(app) as client:
            response = client.get("/health")

        self.assertEqual(response.status_code, 200)
        payload = response.json()
        assert_envelope(self, payload)
        self.assertTrue(payload["success"])
        self.assertEqual(payload["data"], {"status": "ok"})
        self.assertIsNone(payload["error_code"])

    def test_upload_document_is_persisted_saved_and_listed(self):
        with TestClient(app) as client:
            upload = client.post(
                "/upload",
                files={"file": ("course.txt", b"hello database", "text/plain")},
            )
            documents = client.get("/documents")

        upload_payload = upload.json()
        assert_envelope(self, upload_payload)
        self.assertTrue(upload_payload["success"])
        self.assertEqual(upload_payload["data"]["filename"], "course.txt")
        document_id = upload_payload["data"]["document_id"]
        self.assertTrue((Path(settings.upload_dir) / document_id / "course.txt").exists())

        documents_payload = documents.json()
        assert_envelope(self, documents_payload)
        self.assertEqual(len(documents_payload["data"]), 1)
        self.assertEqual(documents_payload["data"][0]["document_id"], document_id)

    def test_upload_rejects_unsupported_and_empty_files(self):
        with TestClient(app) as client:
            xlsx = client.post("/upload", files={"file": ("bad.xlsx", b"abc", "application/octet-stream")})
            doc = client.post("/upload", files={"file": ("bad.doc", b"abc", "application/msword")})
            empty = client.post("/upload", files={"file": ("empty.txt", b"", "text/plain")})

        self.assertEqual(xlsx.json()["error_code"], "UPLOAD_FILE_TYPE_UNSUPPORTED")
        self.assertEqual(doc.json()["error_code"], "UPLOAD_FILE_TYPE_UNSUPPORTED")
        self.assertEqual(empty.json()["error_code"], "UPLOAD_FILE_EMPTY")

    def test_document_persists_across_new_test_client(self):
        with TestClient(app) as client:
            upload = client.post(
                "/upload",
                files={"file": ("persist.txt", b"persistent document", "text/plain")},
            )
            document_id = upload.json()["data"]["document_id"]

        with TestClient(app) as client:
            documents = client.get("/documents")

        document_ids = [item["document_id"] for item in documents.json()["data"]]
        self.assertIn(document_id, document_ids)

    def test_rebuild_processes_txt_chunks_json_faiss_retrieval_and_deduplicates(self):
        with TestClient(app) as client:
            upload = client.post(
                "/upload",
                files={
                    "file": (
                        "rebuild.txt",
                        (
                            "Software development lifecycle includes requirements analysis, "
                            "design, coding, testing, deployment, and maintenance.\n\n"
                            "Agile development emphasizes iterative delivery and feedback."
                        ).encode("utf-8"),
                        "text/plain",
                    )
                },
            )
            document_id = upload.json()["data"]["document_id"]
            first_rebuild = client.post(f"/documents/{document_id}/rebuild")
            second_rebuild = client.post(f"/documents/{document_id}/rebuild")

        first_payload = first_rebuild.json()
        second_payload = second_rebuild.json()
        assert_envelope(self, first_payload)
        self.assertTrue(first_payload["success"])
        self.assertEqual(first_payload["data"]["status"], "processed")
        self.assertGreater(first_payload["data"]["chunk_count"], 0)
        self.assertEqual(second_payload["data"]["chunk_count"], first_payload["data"]["chunk_count"])

        with SessionLocal() as db:
            chunks = db.scalars(select(ChunkModel).where(ChunkModel.document_id == document_id)).all()
        self.assertEqual(len(chunks), first_payload["data"]["chunk_count"])
        self.assertEqual([chunk.chunk_index for chunk in chunks], list(range(1, len(chunks) + 1)))

        self.assertTrue((Path(settings.processed_dir) / f"{document_id}.json").exists())
        self.assertTrue((Path(settings.vector_store_dir) / "index.faiss").exists())
        self.assertTrue((Path(settings.vector_store_dir) / "metadata.json").exists())

        retriever = Retriever(
            index_dir=settings.vector_store_dir,
            embedding_provider=EmbeddingProvider(mode="mock"),
        )
        retrieved_chunks = retriever.retrieve("What does agile development emphasize?", top_k=5)
        self.assertGreater(len(retrieved_chunks), 0)
        self.assertLessEqual(len(retrieved_chunks), 5)
        self.assertEqual(retrieved_chunks[0]["document_id"], document_id)
        self.assertIn("score", retrieved_chunks[0])
        self.assertGreaterEqual(retrieved_chunks[0]["score"], 0)
        self.assertLessEqual(retrieved_chunks[0]["score"], 1)

    def test_rebuild_processes_docx_and_pdf(self):
        cases = [
            ("sample.docx", build_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ("sample.pdf", build_pdf_bytes(), "application/pdf"),
        ]

        for filename, content, media_type in cases:
            with self.subTest(filename=filename):
                with TestClient(app) as client:
                    upload = client.post("/upload", files={"file": (filename, content, media_type)})
                    document_id = upload.json()["data"]["document_id"]
                    rebuild = client.post(f"/documents/{document_id}/rebuild")

                payload = rebuild.json()
                assert_envelope(self, payload)
                self.assertTrue(payload["success"])
                self.assertEqual(payload["data"]["status"], "processed")
                self.assertGreater(payload["data"]["chunk_count"], 0)
                self.assertTrue((Path(settings.upload_dir) / document_id / filename).exists())
                self.assertTrue((Path(settings.processed_dir) / f"{document_id}.json").exists())

    def test_real_embedding_mode_fails_when_local_model_is_missing(self):
        missing_model_path = _tmproot / "models" / "missing-bge"
        with self.assertRaises(FileNotFoundError):
            EmbeddingProvider(model_name=str(missing_model_path), mode="real")

    def test_relative_embedding_model_path_resolves_from_project_root(self):
        provider = EmbeddingProvider(model_name="models/bge-small-zh-v1.5", mode="mock")

        self.assertTrue(Path(provider.model_name).is_absolute())
        self.assertTrue(provider.model_name.endswith(str(Path("models") / "bge-small-zh-v1.5")))

    def test_qa_chain_refuses_when_sources_are_unreliable(self):
        answer = generate_chat_answer(
            question="What is outside the course material?",
            retrieved_chunks=[
                {
                    "chunk_id": "chunk_doc_001_0001",
                    "document_id": "doc_001",
                    "content": "weak source",
                    "score": 0.01,
                    "source": {"filename": "weak.txt", "page": None, "chunk_index": 1},
                }
            ],
            model="deepseek",
        )

        self.assertEqual(answer["sources"], [])
        self.assertIn("无法找到足够可靠", answer["answer"])

    def test_chat_query_returns_vector_index_not_ready_without_faiss(self):
        with TestClient(app) as client:
            chat = client.post("/chat/query", json={"question": "What is the software lifecycle?", "top_k": 5})

        payload = chat.json()
        assert_envelope(self, payload)
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error_code"], "VECTOR_INDEX_NOT_READY")

    def test_chat_query_wraps_missing_embedding_model_error(self):
        with TestClient(app) as client:
            upload = client.post(
                "/upload",
                files={
                    "file": (
                        "missing-model-chat.txt",
                        b"Cloud computing includes measured service and rapid elasticity.",
                        "text/plain",
                    )
                },
            )
            document_id = upload.json()["data"]["document_id"]
            rebuild = client.post(f"/documents/{document_id}/rebuild")
            self.assertTrue(rebuild.json()["success"])

            previous_mode = os.environ.get("RAG_EMBEDDING_MODE")
            previous_path = os.environ.get("RAG_EMBEDDING_MODEL_PATH")
            os.environ["RAG_EMBEDDING_MODE"] = "real"
            os.environ["RAG_EMBEDDING_MODEL_PATH"] = "models/definitely-missing-bge"
            embedding_module._DEFAULT_PROVIDER = None
            try:
                chat = client.post("/chat/query", json={"question": "What are the cloud features?", "top_k": 5})
            finally:
                if previous_mode is None:
                    os.environ.pop("RAG_EMBEDDING_MODE", None)
                else:
                    os.environ["RAG_EMBEDDING_MODE"] = previous_mode
                if previous_path is None:
                    os.environ.pop("RAG_EMBEDDING_MODEL_PATH", None)
                else:
                    os.environ["RAG_EMBEDDING_MODEL_PATH"] = previous_path
                embedding_module._DEFAULT_PROVIDER = None

        payload = chat.json()
        assert_envelope(self, payload)
        self.assertFalse(payload["success"])
        self.assertEqual(payload["error_code"], "QA_CHAIN_FAILED")
        self.assertIn("embedding", payload["message"])

    def test_chat_query_uses_faiss_qa_chain_persists_log_and_updates_stats(self):
        with TestClient(app) as client:
            upload = client.post(
                "/upload",
                files={
                    "file": (
                        "chat.txt",
                        (
                            "Cloud computing has five essential characteristics: on-demand self-service, "
                            "broad network access, resource pooling, rapid elasticity, and measured service."
                        ).encode("utf-8"),
                        "text/plain",
                    )
                },
            )
            document_id = upload.json()["data"]["document_id"]
            rebuild = client.post(f"/documents/{document_id}/rebuild")
            self.assertTrue(rebuild.json()["success"])

            before_stats = client.get("/stats").json()["data"]
            chat = client.post("/chat/query", json={"question": "What are the cloud computing characteristics?", "top_k": 5})
            logs = client.get("/logs")
            after_stats = client.get("/stats").json()["data"]

        chat_payload = chat.json()
        assert_envelope(self, chat_payload)
        self.assertTrue(chat_payload["success"])
        self.assertEqual(chat_payload["data"]["question"], "What are the cloud computing characteristics?")
        self.assertTrue(chat_payload["data"]["answer"])
        self.assertIsInstance(chat_payload["data"]["sources"], list)
        self.assertGreater(len(chat_payload["data"]["sources"]), 0)
        source = chat_payload["data"]["sources"][0]
        self.assertEqual(set(source.keys()), {"filename", "page", "chunk_index", "score", "preview"})
        self.assertTrue(source["preview"])

        logs_payload = logs.json()
        assert_envelope(self, logs_payload)
        self.assertEqual(len(logs_payload["data"]), 1)
        self.assertEqual(logs_payload["data"][0]["question"], "What are the cloud computing characteristics?")

        self.assertEqual(before_stats["question_count"], 0)
        self.assertEqual(after_stats["question_count"], 1)
        self.assertIsNotNone(after_stats["latest_question_time"])

    def test_stats_document_count_changes_after_upload(self):
        with TestClient(app) as client:
            before = client.get("/stats").json()["data"]
            client.post("/upload", files={"file": ("stats.txt", b"stats document", "text/plain")})
            after = client.get("/stats").json()["data"]

        self.assertEqual(before["document_count"], 0)
        self.assertEqual(after["document_count"], 1)


if __name__ == "__main__":
    unittest.main()
